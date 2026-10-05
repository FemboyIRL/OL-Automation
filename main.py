from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from models.models import Applicant, AttorneyType, Case, EvaluationType
from parser.parser import crear_applicant, extraer_elementos
from utils.date import parse_date


def formatear_carpeta_evaluacion(fecha: datetime.date) -> str:
    """Devuelve el nombre de carpeta con formato 'Month Day, Year'."""
    return f"{fecha.strftime('%B')} {fecha.day}, {fecha.year}"


def configurar_estilo_documento(document: Document) -> None:
    """Aplica el estilo base (Times New Roman 11) al documento."""
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # Compatibilidad con Word
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def agregar_attorneys(document: Document, applicant) -> None:
    """Agrega los attorneys en el orden predefinido, si existen."""
    attorney_order = [
        AttorneyType.INSURANCE_COMPANY,
        AttorneyType.APPLICANT,
        AttorneyType.DEFENDANT,
    ]

    for attorney_type in attorney_order:
        attorney = next(
            (a for a in applicant.attorneys if a.type == attorney_type),
            None,
        )
        if attorney is None:
            continue

        document.add_paragraph(str(attorney))
        document.add_paragraph()


def create_folder(evaluation_folder: str, folder_name: str):
    folder = Path("applicants") / evaluation_folder / folder_name
    folder.mkdir(parents=True, exist_ok=True)

    file_path = folder / folder

    return file_path


def crear_word(applicant, evaluation_folder: str) -> Path:
    """Crea el documento Word para un applicant y devuelve su ruta."""

    file_path = create_folder(evaluation_folder, applicant.name)

    document = Document()
    configurar_estilo_documento(document)

    agregar_attorneys(document, applicant)
    document.add_paragraph(applicant.case_information)
    document.add_paragraph(applicant.applicant_description)
    document.add_paragraph(applicant.text_message_to_applicant)

    document.save(file_path)
    return file_path


def cargar_applicants(templates_folder: Path) -> list:
    """Lee los PDFs de templates y construye la lista de applicants."""
    pdf_files = list(templates_folder.glob("*.pdf"))

    if not pdf_files:
        return []

    applicants = []

    for pdf_path in pdf_files:
        print(f"Procesando: {pdf_path.name}")

        elementos = extraer_elementos(pdf_path)
        applicant = crear_applicant(elementos)

        applicants.append(applicant)

    return applicants


def extract_subject(file_path: Path):
    with open(
        file_path,
        "rb",
    ) as f:
        # Analiza el contenido del archivo usando la política por defecto, de momento hay que hacerlo sencillo, pero puede llegar a cambiar por la forma en que nos asignan los casos
        msg = BytesParser(policy=policy.default).parse(f)

        asunto = msg["subject"]

        return asunto


def parse_subject(subject: str) -> Case | None:
    pattern = r"^(?P<type>.+?) - (?P<applicant>.+?) - (?P<date>.+?) -(?P<doctor>[^-]+)(?: - (?P<eval_type>.+))?$"
    match = re.search(pattern, subject)

    if match:
        data = {k: v.strip() if v else None for k, v in match.groupdict().items()}

        eval_date = parse_date(data["date"])

        return Case(
            applicant_name=data["applicant"],
            evaluation_date=eval_date,
            doctor=data["doctor"],
            evaluation_type=EvaluationType.from_str(data["eval_type"]),
        )

    return None


def cargar_correos(email_folder: Path) -> list:
    emails = list(email_folder.glob("*.eml"))

    if not emails:
        return []

    cases = []

    for email in emails:
        subject = extract_subject(email)
        print(f"Procesando: {subject}")

        data = parse_subject(subject)
        if data:
            cases.append(data)
        else:
            print(f"No se pudo procesar el asunto: {subject}")

    return cases


def main() -> None:
    email_folder = Path("emails")
    cases = cargar_correos(email_folder)
    cases.sort(key=lambda a: a.evaluation_date)

    for case in cases:
        evaluation_folder = formatear_carpeta_evaluacion(case.evaluation_date)
        file_path = create_folder(evaluation_folder, case.folder_name)
        print(f"Carpeta creada {file_path}")

    # templates_folder = Path("templates")

    # applicants = cargar_applicants(templates_folder)

    # if not applicants:
    #     print("No se encontraron demos dentro de templates")
    #     return

    # # Ordenar por fecha de evaluación
    # applicants.sort(key=lambda a: a.evaluation_date)

    # for applicant in applicants:
    #     evaluation_folder = formatear_carpeta_evaluacion(applicant.evaluation_date)
    #     file_path = crear_word(applicant, evaluation_folder)
    #     print(f"Word creado: {file_path}")


if __name__ == "__main__":
    main()
